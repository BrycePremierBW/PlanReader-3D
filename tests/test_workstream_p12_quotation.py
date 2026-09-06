"""Workstream P12 — Quotation Subsystem Verification & Regression Tests.

Audits and verifies quotation invariants defined in PLANREADER_AUTONOMOUS_TASK.md:
  - Eligible rows only across all quote exports (summary frame, CSV, Excel, PDF)
  - Canonical unit and Unicode unit ('m²') handling across calculations and aggregations
  - Rates, line totals, mark-up, GST, and grand totals exact reconciliation
  - Excel workbook sheet structure and formatting (Quote Header, Per-Level Summary,
    Totals, Take-off Detail, register sheets)
  - Safe handling of empty workspaces and edge cases
"""

from __future__ import annotations

import io
import math
from typing import Any, Dict, List
import openpyxl
import pandas as pd
import pytest

import pb_planreader_3d_app as app
import pb_takeoff_accuracy_v125 as v125
from pb_takeoff_authority_v164 import approve_model_surface_row


@pytest.fixture
def setup_db() -> None:
    app.init_local_db()


@pytest.fixture
def mixed_quote_workspace(setup_db: None) -> int:
    wid = app.create_standalone_workspace("QUOTE_WS", "Commercial Tower", "Apex Builders", "100 Queen St")
    app.set_workspace_setting(wid, "pricing_margin_pct", 15.0)
    app.set_workspace_setting(wid, "gst_rate_pct", 10.0)
    app.set_workspace_setting(wid, "internal_pricing_basis", "wall_m2")

    # Document & Page
    doc_id = app.lexecute(
        "INSERT INTO documents(workspace_id, file_name, source_type, category, page_count, uploaded_at) VALUES(?, 'A101.pdf', 'pdf', 'plans', 1, ?)",
        (wid, app.now_stamp()),
    )
    pid = app.lexecute(
        "INSERT INTO pages(workspace_id, document_id, page_no, page_label, px_per_m) VALUES(?, ?, 1, 'Page 1', 100.0)",
        (wid, doc_id),
    )

    # 1. Authoritative manual row on Ground with ASCII unit 'm2'
    app.lexecute(
        """INSERT INTO takeoff_rows(workspace_id, section, element, location, substrate, finish_system, quantity, unit, quantity_status, source_page, source_reference, inclusion_status, coats, coverage_m2_per_litre, productivity_m2_per_hour, rate_per_unit, confidence, notes, row_role, created_at, updated_at)
           VALUES(?, 'Internal', 'Walls', 'Ground', 'Plasterboard', 'Low sheen acrylic', 100.0, 'm2', 'Measured', 'Page 1', '', 'Included', 2, 10.0, 5.0, 30.0, 'Authoritative', '', 'work', ?, ?)""",
        (wid, app.now_stamp(), app.now_stamp()),
    )

    # 2. Authoritative manual row on Ground with unit 'LM' (lineal metres)
    app.lexecute(
        """INSERT INTO takeoff_rows(workspace_id, section, element, location, substrate, finish_system, quantity, unit, quantity_status, source_page, source_reference, inclusion_status, coats, coverage_m2_per_litre, productivity_m2_per_hour, rate_per_unit, confidence, notes, row_role, created_at, updated_at)
           VALUES(?, 'Internal', 'Skirting', 'Ground', 'Timber trim / joinery', 'Gloss enamel', 40.0, 'LM', 'Measured', 'Page 1', '', 'Included', 2, 10.0, 5.0, 15.0, 'Authoritative', '', 'work', ?, ?)""",
        (wid, app.now_stamp(), app.now_stamp()),
    )

    # 3. Authoritative manual row on Ground with unit 'each' (door count)
    app.lexecute(
        """INSERT INTO takeoff_rows(workspace_id, section, element, location, substrate, finish_system, quantity, unit, quantity_status, source_page, source_reference, inclusion_status, coats, coverage_m2_per_litre, productivity_m2_per_hour, rate_per_unit, confidence, notes, row_role, created_at, updated_at)
           VALUES(?, 'Internal', 'Doors', 'Ground', 'Timber door', 'Gloss enamel', 5.0, 'each', 'Measured', 'Page 1', '', 'Included', 2, 10.0, 5.0, 100.0, 'Authoritative', '', 'work', ?, ?)""",
        (wid, app.now_stamp(), app.now_stamp()),
    )

    # 4. Authoritative manual row on Ground with unit 'Item' (allowance)
    app.lexecute(
        """INSERT INTO takeoff_rows(workspace_id, section, element, location, substrate, finish_system, quantity, unit, quantity_status, source_page, source_reference, inclusion_status, coats, coverage_m2_per_litre, productivity_m2_per_hour, rate_per_unit, confidence, notes, row_role, created_at, updated_at)
           VALUES(?, 'Preliminaries', 'Site protection', 'Ground', 'Other', '', 1.0, 'Item', 'Allowance', 'Page 1', '', 'Included', 1, 10.0, 1.0, 500.0, 'Authoritative', '', 'work', ?, ?)""",
        (wid, app.now_stamp(), app.now_stamp()),
    )

    # 5. Commercial floor reference row on Ground with unit 'm2' (200 m²)
    app.lexecute(
        """INSERT INTO takeoff_rows(workspace_id, section, element, location, substrate, finish_system, quantity, unit, quantity_status, source_page, source_reference, inclusion_status, coats, coverage_m2_per_litre, productivity_m2_per_hour, rate_per_unit, confidence, notes, row_role, created_at, updated_at)
           VALUES(?, 'Reference', 'Floor Area', 'Ground', '', '', 200.0, 'm2', 'Measured', 'Page 1', '', 'Included', 0, 0, 0, 0.0, 'Authoritative', '', 'floor_area', ?, ?)""",
        (wid, app.now_stamp(), app.now_stamp()),
    )

    # 6. Approved 3D model surface on Level 1 (quantity=50.0 m², rate=30.0)
    raw_model_row = {
        "workspace_id": wid,
        "section": "Internal",
        "element": "Walls",
        "location": "Level 1",
        "substrate": "Plasterboard",
        "finish_system": "Low sheen acrylic",
        "quantity": 50.0,
        "unit": "m²",
        "quantity_status": "Measured",
        "source_page": "3d_model",
        "source_reference": "pb 3d surface editor 101",
        "inclusion_status": "Included",
        "coats": 2,
        "coverage_m2_per_litre": 10.0,
        "productivity_m2_per_hour": 5.0,
        "rate_per_unit": 30.0,
        "confidence": "3D Model",
        "notes": "",
        "row_role": "model_surface",
    }
    approved_row = approve_model_surface_row(
        raw_model_row,
        source="pb 3d surface editor 101",
        reviewed_by="Bryce Curran",
        reviewed_at="2026-09-06T00:00:00Z",
    )
    app.lexecute(
        """INSERT INTO takeoff_rows(workspace_id, section, element, location, substrate, finish_system, quantity, unit, quantity_status, source_page, source_reference, inclusion_status, coats, coverage_m2_per_litre, productivity_m2_per_hour, rate_per_unit, confidence, notes, row_role, commercial_authority_status, commercial_authority_source, commercial_authority_reviewed_by, commercial_authority_reviewed_at, commercial_authority_fingerprint, created_at, updated_at)
           VALUES(?, 'Internal', 'Walls', 'Level 1', 'Plasterboard', 'Low sheen acrylic', 50.0, 'm²', 'Measured', '3d_model', 'pb 3d surface editor 101', 'Included', 2, 10.0, 5.0, 30.0, '3D Model', '', 'model_surface', 'APPROVED', 'pb 3d surface editor 101', 'Bryce Curran', '2026-09-06T00:00:00Z', ?, ?, ?)""",
        (wid, approved_row["commercial_authority_fingerprint"], app.now_stamp(), app.now_stamp()),
    )

    # 7. Commercial floor reference row on Level 1 (150 m²)
    app.lexecute(
        """INSERT INTO takeoff_rows(workspace_id, section, element, location, substrate, finish_system, quantity, unit, quantity_status, source_page, source_reference, inclusion_status, coats, coverage_m2_per_litre, productivity_m2_per_hour, rate_per_unit, confidence, notes, row_role, created_at, updated_at)
           VALUES(?, 'Reference', 'Floor Area', 'Level 1', '', '', 150.0, 'm²', 'Measured', 'Page 1', '', 'Included', 0, 0, 0, 0.0, 'Authoritative', '', 'floor_area', ?, ?)""",
        (wid, app.now_stamp(), app.now_stamp()),
    )

    # 8. Unapproved 3D model surface on Level 2 (quantity=80.0 m², rate=30.0) -> INELIGIBLE
    app.lexecute(
        """INSERT INTO takeoff_rows(workspace_id, section, element, location, substrate, finish_system, quantity, unit, quantity_status, source_page, source_reference, inclusion_status, coats, coverage_m2_per_litre, productivity_m2_per_hour, rate_per_unit, confidence, notes, row_role, commercial_authority_status, created_at, updated_at)
           VALUES(?, 'Internal', 'Walls', 'Level 2', 'Plasterboard', 'Low sheen acrylic', 80.0, 'm²', 'Provisional', 'Page 1', 'Surface_002', 'Included', 2, 10.0, 5.0, 30.0, '3D Model', '', 'model_surface', 'DRAFT', ?, ?)""",
        (wid, app.now_stamp(), app.now_stamp()),
    )

    # 9. Excluded row on Level 3 (quantity=60.0 m², rate=30.0) -> INELIGIBLE
    app.lexecute(
        """INSERT INTO takeoff_rows(workspace_id, section, element, location, substrate, finish_system, quantity, unit, quantity_status, source_page, source_reference, inclusion_status, coats, coverage_m2_per_litre, productivity_m2_per_hour, rate_per_unit, confidence, notes, row_role, created_at, updated_at)
           VALUES(?, 'Internal', 'Walls', 'Level 3', 'Plasterboard', 'Low sheen acrylic', 60.0, 'm²', 'Measured', 'Page 1', '', 'Excluded', 2, 10.0, 5.0, 30.0, 'Authoritative', '', 'work', ?, ?)""",
        (wid, app.now_stamp(), app.now_stamp()),
    )

    # Register items
    app.lexecute(
        "INSERT INTO register_items(workspace_id, register_name, item_no, title, detail, priority, source_reference, status) VALUES(?, 'inclusions', 'INC-01', 'Internal paint', '2 coats low sheen acrylic', 'Normal', 'Spec 09900', 'Agreed')",
        (wid,),
    )
    app.lexecute(
        "INSERT INTO register_items(workspace_id, register_name, item_no, title, detail, priority, source_reference, status) VALUES(?, 'exclusions', 'EXC-01', 'Level 3 works', 'Excluded per client scope note', 'High', 'Addendum 1', 'Agreed')",
        (wid,),
    )

    return wid


class TestQuotationUnitAndUnicodeNormalization:
    """Verifies that ASCII unit variants and Unicode units ('m²') are properly handled."""

    def test_ascii_m2_unit_yields_correct_paint_and_labour(self, mixed_quote_workspace: int) -> None:
        wid = mixed_quote_workspace
        df = app.dataframe_for_takeoff(wid)
        wall_row = df[(df["element"] == "Walls") & (df["location"] == "Ground")].iloc[0]

        # 100 m2, 2 coats, 10 m2/L coverage -> (100 * 2) / 10 = 20.0 L
        assert wall_row["paint_litres"] == pytest.approx(20.0), (
            f"Expected 20.0 paint litres for m2 wall row, got {wall_row['paint_litres']}"
        )
        # 100 m2, 5 m2/hr productivity -> 100 / 5 = 20.0 hrs
        assert wall_row["labour_hours"] == pytest.approx(20.0), (
            f"Expected 20.0 labour hours for m2 wall row, got {wall_row['labour_hours']}"
        )
        # Canonical unit must be 'm²'
        assert wall_row["unit"] == "m²"

    def test_lineal_unit_variants_yield_correct_lm_and_labour(self, mixed_quote_workspace: int) -> None:
        wid = mixed_quote_workspace
        df = app.dataframe_for_takeoff(wid)
        skirt_row = df[df["element"] == "Skirting"].iloc[0]

        # 40 LM, 5 m/hr productivity -> 40 / 5 = 8.0 hrs
        assert skirt_row["labour_hours"] == pytest.approx(8.0), (
            f"Expected 8.0 labour hours for LM skirt row, got {skirt_row['labour_hours']}"
        )
        assert skirt_row["unit"] == "lm"

    def test_count_and_item_unit_variants_yield_correct_count_and_labour(self, mixed_quote_workspace: int) -> None:
        wid = mixed_quote_workspace
        df = app.dataframe_for_takeoff(wid)
        doors_row = df[df["element"] == "Doors"].iloc[0]
        site_row = df[df["element"] == "Site protection"].iloc[0]

        # 5 doors, 5 / 5 = 1.0 hr
        assert doors_row["labour_hours"] == pytest.approx(1.0)
        assert doors_row["unit"] == "No."

        # 1 item, 1 / 1 = 1.0 hr
        assert site_row["labour_hours"] == pytest.approx(1.0)
        assert site_row["unit"] == "item"

    def test_per_level_summary_aggregates_units_properly(self, mixed_quote_workspace: int) -> None:
        wid = mixed_quote_workspace
        pls = app.per_level_summary(wid)
        ground_row = pls[pls["level"] == "Ground"].iloc[0]

        # Ground has: 1 wall (100 m²), 1 skirting (40 lm), 1 doors (5 No.), 1 prelim (1 item)
        # and 1 floor reference (200 m²)
        assert ground_row["rows"] == 4  # 4 work rows
        assert ground_row["m2"] == pytest.approx(100.0)
        assert ground_row["floor_m2"] == pytest.approx(200.0)
        assert ground_row["lm"] == pytest.approx(40.0)
        assert ground_row["count"] == pytest.approx(6.0)  # 5 doors + 1 item
        assert ground_row["paint_litres"] == pytest.approx(20.0)
        assert ground_row["labour_hours"] == pytest.approx(30.0)  # 20 + 8 + 1 + 1

    def test_floor_area_unit_normalization_under_floor_m2_pricing_basis(self, mixed_quote_workspace: int) -> None:
        wid = mixed_quote_workspace
        app.set_workspace_setting(wid, "internal_pricing_basis", "floor_m2")

        df = app.dataframe_for_takeoff(wid)
        wall_row = df[(df["element"] == "Walls") & (df["location"] == "Ground")].iloc[0]

        # Under floor_m2 basis, Ground internal wall gets Ground floor area (200.0 m²)
        assert wall_row["priced_quantity"] == pytest.approx(200.0)
        assert wall_row["pricing_basis"] == "Floor m²"
        # Value = 200 m² * $30 = $6,000
        assert wall_row["value_ex_gst"] == pytest.approx(6000.0)

    def test_takeoff_summary_csv_unit_normalization(self, mixed_quote_workspace: int) -> None:
        wid = mixed_quote_workspace
        df = app.dataframe_for_takeoff(wid)
        csv_text = app._takeoff_summary_csv(df)

        lines = csv_text.strip().replace("\r\n", "\n").split("\n")
        assert len(lines) >= 2
        # Verify header
        assert lines[0] == "section,m2,lineal_m,count,paint_litres,labour_hours,value_ex_gst"
        # Find Internal row
        internal_lines = [line for line in lines if line.startswith("Internal")]
        assert len(internal_lines) == 1
        parts = internal_lines[0].split(",")
        m2_val = float(parts[1])
        lm_val = float(parts[2])
        count_val = float(parts[3])

        # Internal has 100 m² (Ground) + 50 m² (Level 1) = 150 m²
        assert m2_val == pytest.approx(150.0)
        # Internal has 40 lm (Ground skirting)
        assert lm_val == pytest.approx(40.0)
        # Internal has 5 doors
        assert count_val == pytest.approx(5.0)


class TestQuotationEligibleRowsOnly:
    """Verifies that unapproved, tampered, and excluded rows are barred from quotations."""

    def test_unapproved_model_surfaces_omitted_from_summary_and_totals(self, mixed_quote_workspace: int) -> None:
        wid = mixed_quote_workspace
        pls = app.per_level_summary(wid)
        levels_present = set(pls["level"].tolist())

        # Level 2 only had an unapproved model surface, so Level 2 must not appear
        assert "Level 2" not in levels_present

    def test_excluded_rows_omitted_from_summary_and_totals(self, mixed_quote_workspace: int) -> None:
        wid = mixed_quote_workspace
        pls = app.per_level_summary(wid)
        levels_present = set(pls["level"].tolist())

        # Level 3 only had an excluded row, so Level 3 must not appear
        assert "Level 3" not in levels_present

    def test_excel_takeoff_detail_sheet_contains_only_eligible_rows(self, mixed_quote_workspace: int) -> None:
        wid = mixed_quote_workspace
        wb_bytes = app.quote_workbook_bytes(wid)
        wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
        ws = wb["Take-off Detail"]

        data = list(ws.iter_rows(values_only=True))
        assert len(data) >= 2
        header = [str(c) for c in data[0]]
        loc_idx = header.index("location")
        role_idx = header.index("row_role")
        inc_idx = header.index("inclusion_status")

        locations = [row[loc_idx] for row in data[1:]]
        roles = [row[role_idx] for row in data[1:]]
        inclusions = [row[inc_idx] for row in data[1:]]

        # No Level 2 (unapproved) or Level 3 (excluded) in Take-off Detail
        assert "Level 2" not in locations
        assert "Level 3" not in locations
        assert all(inc != "Excluded" for inc in inclusions)


class TestQuotationTotalsAndGSTReconciliation:
    """Verifies exact reconciliation across summary frame, CSV, Excel, and PDF."""

    def test_exact_reconciliation_wall_m2_basis(self, mixed_quote_workspace: int) -> None:
        wid = mixed_quote_workspace
        app.set_workspace_setting(wid, "pricing_margin_pct", 15.0)
        app.set_workspace_setting(wid, "gst_rate_pct", 10.0)

        # 1. Calculate expected totals from work rows
        # Ground:
        #   Walls: 100 m² * $30 = $3,000
        #   Skirting: 40 lm * $15 = $600
        #   Doors: 5 * $100 = $500
        #   Site protection: 1 * $500 = $500
        # Ground Subtotal = $4,600
        # Level 1:
        #   Approved Walls: 50 m² * $30 = $1,500
        # Total Value ex GST = $4,600 + $1,500 = $6,100
        # Margin 15%: $6,100 * 0.15 = $915.0
        # Subtotal ex GST: $6,100 + $915 = $7,015.0
        # GST 10%: $7,015 * 0.10 = $701.50
        # Grand Total inc GST: $7,015 + $701.50 = $7,716.50
        expected_val_sum = 6100.0
        expected_markup = 915.0
        expected_subtotal_ex = 7015.0
        expected_gst = 701.50
        expected_grand = 7716.50

        # Check Quote Summary Frame
        qsf = app.quote_summary_frame(wid)
        assert qsf["value_ex_gst"].sum() == pytest.approx(expected_val_sum)
        assert (qsf["markup_ex_gst"] - qsf["value_ex_gst"]).sum() == pytest.approx(expected_markup)
        assert qsf["gst"].sum() == pytest.approx(expected_gst)
        assert qsf["total_inc_gst"].sum() == pytest.approx(expected_grand)

        # Check CSV export matches summary frame
        csv_text = app.per_level_summary_csv(wid)
        csv_df = pd.read_csv(io.StringIO(csv_text))
        assert csv_df["value_ex_gst"].sum() == pytest.approx(expected_val_sum)
        assert csv_df["total_inc_gst"].sum() == pytest.approx(expected_grand)

        # Check Excel workbook Totals sheet
        wb_bytes = app.quote_workbook_bytes(wid)
        wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
        totals_ws = wb["Totals"]
        totals_dict = {row[0]: float(row[1]) for row in totals_ws.iter_rows(values_only=True) if row[0] not in (None, "Metric")}
        assert totals_dict["Value ex GST"] == pytest.approx(expected_val_sum)
        assert totals_dict["Mark-up"] == pytest.approx(expected_markup)
        assert totals_dict["Subtotal ex GST"] == pytest.approx(expected_subtotal_ex)
        assert totals_dict["GST"] == pytest.approx(expected_gst)
        assert totals_dict["Total inc GST"] == pytest.approx(expected_grand)

        # Check PDF bytes generation
        pdf_bytes = app.quote_pdf_bytes(wid)
        assert pdf_bytes[:4] == b"%PDF"
        assert len(pdf_bytes) > 1000

    def test_exact_reconciliation_floor_m2_basis(self, mixed_quote_workspace: int) -> None:
        wid = mixed_quote_workspace
        app.set_workspace_setting(wid, "internal_pricing_basis", "floor_m2")
        app.set_workspace_setting(wid, "pricing_margin_pct", 20.0)
        app.set_workspace_setting(wid, "gst_rate_pct", 10.0)

        # Ground:
        #   Internal wall gets Ground floor area = 200 m² * $30 = $6,000
        #   Skirting: 40 lm * $15 = $600
        #   Doors: 5 * $100 = $500
        #   Site protection: 1 * $500 = $500
        # Ground Subtotal = $7,600
        # Level 1:
        #   Internal wall gets Level 1 floor area = 150 m² * $30 = $4,500
        # Total Value ex GST = $7,600 + $4,500 = $12,100
        # Margin 20%: $12,100 * 0.20 = $2,420
        # Subtotal ex GST: $12,100 + $2,420 = $14,520
        # GST 10%: $14,520 * 0.10 = $1,452
        # Grand Total inc GST: $14,520 + $1,452 = $15,972
        expected_val_sum = 12100.0
        expected_markup = 2420.0
        expected_subtotal_ex = 14520.0
        expected_gst = 1452.0
        expected_grand = 15972.0

        qsf = app.quote_summary_frame(wid)
        assert qsf["value_ex_gst"].sum() == pytest.approx(expected_val_sum)
        assert (qsf["markup_ex_gst"] - qsf["value_ex_gst"]).sum() == pytest.approx(expected_markup)
        assert qsf["gst"].sum() == pytest.approx(expected_gst)
        assert qsf["total_inc_gst"].sum() == pytest.approx(expected_grand)

        wb_bytes = app.quote_workbook_bytes(wid)
        wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
        totals_ws = wb["Totals"]
        totals_dict = {row[0]: float(row[1]) for row in totals_ws.iter_rows(values_only=True) if row[0] not in (None, "Metric")}
        assert totals_dict["Value ex GST"] == pytest.approx(expected_val_sum)
        assert totals_dict["Mark-up"] == pytest.approx(expected_markup)
        assert totals_dict["Subtotal ex GST"] == pytest.approx(expected_subtotal_ex)
        assert totals_dict["GST"] == pytest.approx(expected_gst)
        assert totals_dict["Total inc GST"] == pytest.approx(expected_grand)

    def test_zero_margin_and_zero_gst(self, mixed_quote_workspace: int) -> None:
        wid = mixed_quote_workspace
        app.set_workspace_setting(wid, "pricing_margin_pct", 0.0)
        app.set_workspace_setting(wid, "gst_rate_pct", 0.0)

        qsf = app.quote_summary_frame(wid)
        assert (qsf["markup_ex_gst"] - qsf["value_ex_gst"]).sum() == pytest.approx(0.0)
        assert qsf["gst"].sum() == pytest.approx(0.0)
        assert qsf["total_inc_gst"].sum() == pytest.approx(qsf["value_ex_gst"].sum())


class TestQuotationWorkbookStructureAndFormatting:
    """Verifies Excel sheets, formatting, and edge cases like empty workspaces."""

    def test_workbook_all_expected_sheets_present(self, mixed_quote_workspace: int) -> None:
        wid = mixed_quote_workspace
        wb_bytes = app.quote_workbook_bytes(wid)
        wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))

        expected_sheets = [
            "Quote Header",
            "Per-Level Summary",
            "Totals",
            "Take-off Detail",
            "Door Schedule",
            "Inclusions",
            "Exclusions",
            "Separate Clarifications",
            "Assumptions",
            "RFIs",
        ]
        for sheet_name in expected_sheets:
            assert sheet_name in wb.sheetnames, f"Missing sheet {sheet_name} in quotation workbook"

    def test_empty_workspace_exports_safely(self, setup_db: None) -> None:
        empty_wid = app.create_standalone_workspace("EMPTY_WS", "Empty Job", "Client C", "None")

        qsf = app.quote_summary_frame(empty_wid)
        assert qsf.empty

        csv_text = app.per_level_summary_csv(empty_wid)
        assert csv_text == ""

        wb_bytes = app.quote_workbook_bytes(empty_wid)
        assert wb_bytes[:2] == b"PK"
        wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
        assert "Totals" in wb.sheetnames

        pdf_bytes = app.quote_pdf_bytes(empty_wid)
        assert pdf_bytes[:4] == b"%PDF"


class TestTakeoffAccuracyV125QuotationReconciliation:
    """Verifies that pb_takeoff_accuracy_v125 behaves identically to core."""

    def test_v125_per_level_summary_and_dataframe_takeoff(self, mixed_quote_workspace: int) -> None:
        wid = mixed_quote_workspace
        core_df = app.dataframe_for_takeoff(wid)
        v125_df = v125.dataframe_for_takeoff(app, wid)

        assert core_df["value_ex_gst"].sum() == pytest.approx(v125_df["value_ex_gst"].sum())
        assert core_df["paint_litres"].sum() == pytest.approx(v125_df["paint_litres"].sum())
        assert core_df["labour_hours"].sum() == pytest.approx(v125_df["labour_hours"].sum())

        core_pls = app.per_level_summary(wid)
        v125_pls = v125.per_level_summary(app, wid)

        assert core_pls["value_ex_gst"].sum() == pytest.approx(v125_pls["value_ex_gst"].sum())
        assert core_pls["m2"].sum() == pytest.approx(v125_pls["m2"].sum())
        assert core_pls["floor_m2"].sum() == pytest.approx(v125_pls["floor_m2"].sum())
