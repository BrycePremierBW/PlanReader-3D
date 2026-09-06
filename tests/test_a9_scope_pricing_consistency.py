"""Focused A9 regressions for unresolved scope in pricing and quote outputs."""

from __future__ import annotations

import io
import sqlite3

import openpyxl
import pytest

import pb_planreader_3d_app as app
from pb_takeoff_authority_v164 import (
    is_jobhub_eligible_row,
    takeoff_row_pricing_authority,
    takeoff_row_publishability,
)


@pytest.fixture
def scope_pricing_workspace(tmp_path, monkeypatch) -> int:
    db_path = str(tmp_path / "a9_scope_pricing.db")

    def _connect() -> sqlite3.Connection:
        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row
        return conn

    monkeypatch.setattr(app, "local_connect", _connect)
    app.init_local_db()
    workspace_id = app.create_standalone_workspace(
        "A9-SCOPE", "A9 scope pricing", "PB", "100 Commercial St"
    )

    base = {
        "workspace_id": workspace_id,
        "section": "Internal",
        "element": "Walls",
        "location": "Level 1",
        "substrate": "Plasterboard",
        "finish_system": "Low sheen acrylic",
        "unit": "m²",
        "source_page": "A-101",
        "source_reference": "Drawing A-101",
        "coats": 2,
        "coverage_m2_per_litre": 10.0,
        "productivity_m2_per_hour": 5.0,
        "rate_per_unit": 10.0,
        "confidence": "Reviewed",
        "notes": "",
        "row_role": "work",
    }
    rows = [
        {**base, "id": 1, "quantity": 10.0, "quantity_status": "Measured", "inclusion_status": "INCLUSION"},
        {**base, "id": 2, "quantity": 20.0, "quantity_status": "Measured", "inclusion_status": "EXCLUSION"},
        {**base, "id": 3, "quantity": 30.0, "quantity_status": "Measured", "inclusion_status": "CLARIFICATION"},
        {**base, "id": 4, "quantity": 40.0, "quantity_status": "Measured", "inclusion_status": "PROVISIONAL"},
        {**base, "id": 5, "quantity": 50.0, "quantity_status": "Provisional measured", "inclusion_status": "INCLUSION"},
        {**base, "id": 6, "quantity": 60.0, "quantity_status": "To measure", "inclusion_status": "INCLUSION"},
        {**base, "id": 7, "quantity": 70.0, "quantity_status": "Excluded", "inclusion_status": "INCLUSION"},
        {**base, "id": 8, "element": "Floor area", "quantity": 10.0, "quantity_status": "Measured", "inclusion_status": "INCLUSION", "rate_per_unit": 0.0, "row_role": "floor_area"},
        {**base, "id": 9, "element": "Floor area", "quantity": 900.0, "quantity_status": "Measured", "inclusion_status": "PROVISIONAL", "rate_per_unit": 0.0, "row_role": "floor_area"},
    ]

    with _connect() as conn:
        for row in rows:
            columns = list(row)
            conn.execute(
                f"INSERT INTO takeoff_rows({','.join(columns)}) "
                f"VALUES({','.join('?' for _ in columns)})",
                [row[column] for column in columns],
            )

    app.set_workspace_setting(workspace_id, "internal_pricing_basis", "floor_m2")
    return workspace_id


@pytest.mark.parametrize(
    ("row", "reason"),
    [
        ({"quantity": 30.0, "quantity_status": "Measured", "inclusion_status": "CLARIFICATION"}, "CLARIFICATION"),
        ({"quantity": 40.0, "quantity_status": "Measured", "inclusion_status": "PROVISIONAL"}, "PROVISIONAL"),
        ({"quantity": 50.0, "quantity_status": "Provisional measured", "inclusion_status": "INCLUSION"}, "PROVISIONAL"),
        ({"quantity": 60.0, "quantity_status": "To measure", "inclusion_status": "INCLUSION"}, "PROVISIONAL"),
        ({"quantity": 70.0, "quantity_status": "Excluded", "inclusion_status": "INCLUSION"}, "EXCLUDED"),
    ],
)
def test_unresolved_scope_is_not_firm_pricing_or_active_work(row: dict, reason: str) -> None:
    assert takeoff_row_pricing_authority(row) == (False, reason)
    assert is_jobhub_eligible_row(row) == (False, reason)


@pytest.mark.parametrize(
    "row",
    [
        {"quantity": 30.0, "quantity_status": "Measured", "inclusion_status": "CLARIFICATION"},
        {"quantity": 40.0, "quantity_status": "Measured", "inclusion_status": "PROVISIONAL"},
        {"quantity": 50.0, "quantity_status": "Provisional measured", "inclusion_status": "INCLUSION"},
    ],
)
def test_unresolved_scope_remains_visible_to_preflight_review(row: dict) -> None:
    assert takeoff_row_publishability(row) == (True, "PUBLISHABLE")


def test_unresolved_scope_cannot_inflate_pricing_or_quote_totals(
    scope_pricing_workspace: int,
) -> None:
    workspace_id = scope_pricing_workspace
    raw = app.ldf(
        "SELECT * FROM takeoff_rows WHERE workspace_id=? ORDER BY id",
        (workspace_id,),
    )

    assert raw["quantity"].sum() == pytest.approx(1190.0)
    assert app.takeoff_work_rows(raw)["id"].tolist() == [1]

    priced = app.dataframe_for_takeoff(workspace_id)
    assert priced["id"].tolist() == [1, 8]
    assert priced.loc[priced["id"].eq(1), "priced_quantity"].item() == pytest.approx(10.0)
    assert priced.loc[priced["id"].eq(8), "value_ex_gst"].item() == 0.0
    assert priced["value_ex_gst"].sum() == pytest.approx(100.0)

    quote = app.quote_summary_frame(workspace_id)
    assert quote["value_ex_gst"].sum() == pytest.approx(100.0)

    workbook = openpyxl.load_workbook(
        io.BytesIO(app.quote_workbook_bytes(workspace_id)), data_only=True
    )
    totals = {
        row[0]: float(row[1])
        for row in workbook["Totals"].iter_rows(values_only=True)
        if row[0] not in (None, "Metric")
    }
    assert totals["Value ex GST"] == pytest.approx(100.0)
    detail = list(workbook["Take-off Detail"].iter_rows(values_only=True))
    id_index = detail[0].index("id")
    assert [row[id_index] for row in detail[1:]] == [1, 8]
